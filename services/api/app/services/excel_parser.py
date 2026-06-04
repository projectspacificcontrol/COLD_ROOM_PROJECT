from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import BinaryIO, Any

from openpyxl import load_workbook


def parse_excel_file(path: str | Path) -> list[dict[str, Any]]:
    # Some vendor exports report an incorrect read-only dimension, hiding wide headers.
    workbook = load_workbook(path, read_only=False, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(header).strip() for header in next(rows) if header is not None]
    records: list[dict[str, Any]] = []
    for row in rows:
        record = {headers[index]: value for index, value in enumerate(row[: len(headers)])}
        if record.get("LogTime"):
            records.append(record)
    return records


def parse_excel_upload(file_obj: BinaryIO) -> list[dict[str, Any]]:
    with NamedTemporaryFile(suffix=".xlsx", delete=False) as temp_file:
        temp_file.write(file_obj.read())
        temp_path = temp_file.name
    return parse_excel_file(temp_path)

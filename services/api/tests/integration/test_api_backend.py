import asyncio
from datetime import datetime, timezone
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import AsyncSessionLocal
from app.main import app
from app.models.entities import AuditLog, IpAllowlistEntry, RawIngestLog, SensorFault
from app.core.config import settings


def test_auth_login_and_me(client: TestClient) -> None:
    csrf = client.get("/api/auth/csrf").json()["csrf_token"]
    login = client.post(
        "/api/auth/login",
        json={"email": "admin@example.com", "password": "ValidPassword123!"},
        headers={"X-CSRF-Token": csrf},
    )
    assert login.status_code == 200
    me = client.get("/api/auth/me")
    assert me.status_code == 200
    assert me.json()["role"] == "admin"

    async def audit_actions():
        async with AsyncSessionLocal() as session:
            return [log.action for log in (await session.execute(select(AuditLog))).scalars().all()]

    assert "login_success" in asyncio.run(audit_actions())


def test_health_ready_and_metrics(client: TestClient) -> None:
    assert client.get("/api/health").json() == {"status": "ok"}
    assert client.get("/api/ready").json() == {"status": "ready"}
    metrics = client.get("/api/metrics")
    assert metrics.status_code == 200
    assert "cold_room_active_alerts" in metrics.text
    assert "cold_room_sensor_faults" in metrics.text


def test_ingestion_handles_missing_sensor_columns(admin_client: TestClient) -> None:
    payload = [{"LogTime": datetime.now(timezone.utc).isoformat(), "CR123_TT01": "4.2"}]
    response = admin_client.post("/api/ingest/records", json=payload)
    assert response.status_code == 200
    body = response.json()
    assert body["rows_inserted"] == 1
    assert body["readings_normalized"] == 90
    assert "CR161718_TT15" in body["missing_expected_sensors"]

    async def read_counts():
        async with AsyncSessionLocal() as session:
            raw_count = len((await session.execute(select(RawIngestLog))).scalars().all())
            fault_count = len((await session.execute(select(SensorFault).where(SensorFault.status == "open"))).scalars().all())
            return raw_count, fault_count

    raw_count, fault_count = asyncio.run(read_counts())
    assert raw_count == 1
    assert fault_count >= 1


def test_report_export_returns_xlsx(admin_client: TestClient) -> None:
    response = admin_client.get("/api/admin/reports/export?duration=1h&format=xlsx&include_alerts=true&include_faults=true")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument")
    assert response.content[:2] == b"PK"
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert workbook.sheetnames == ["Summary", "Room Latest Snapshot", "Temperature Readings", "Alerts", "Sensor Faults", "Metadata"]


def test_ip_allowlist_blocks_unlisted_client() -> None:
    async def add_allowlist():
        async with AsyncSessionLocal() as session:
            session.add(IpAllowlistEntry(cidr="203.0.113.0/24", description="test network"))
            await session.commit()

    asyncio.run(add_allowlist())
    settings.api_trusted_proxy_enabled = True
    blocked = TestClient(app)
    allowed = TestClient(app)

    assert blocked.get("/api/dashboard/overview", headers={"X-Forwarded-For": "198.51.100.5"}).status_code == 403
    assert allowed.get("/api/dashboard/overview", headers={"X-Forwarded-For": "203.0.113.10"}).status_code == 200
    settings.api_trusted_proxy_enabled = False


def test_threshold_priority_and_status_effect(admin_client: TestClient) -> None:
    rooms = admin_client.get("/api/admin/rooms").json()
    room_one = next(room for room in rooms if room["room_number"] == 1)
    sensor_one = room_one["sensors"][0]

    assert admin_client.post("/api/admin/thresholds", json={"scope_type": "group", "scope_key": "CR123", "warning_max_c": 4.0, "critical_max_c": 6.0, "stale_after_seconds": 900}).status_code == 200
    assert admin_client.post("/api/admin/thresholds", json={"scope_type": "room", "scope_id": room_one["id"], "warning_max_c": 3.0, "critical_max_c": 4.0, "stale_after_seconds": 900}).status_code == 200
    assert admin_client.post("/api/admin/thresholds", json={"scope_type": "sensor", "scope_id": sensor_one["id"], "warning_max_c": 2.0, "critical_max_c": 3.0, "stale_after_seconds": 900}).status_code == 200

    effective = admin_client.get(f"/api/admin/thresholds/effective?sensor_id={sensor_one['id']}").json()["rule"]
    assert effective["scope_type"] == "sensor"

    payload = [{"LogTime": datetime.now(timezone.utc).isoformat(), "CR123_TT01": "3.5"}]
    response = admin_client.post("/api/ingest/records", json=payload)
    assert response.status_code == 200

    room = admin_client.get(f"/api/dashboard/rooms/{room_one['id']}").json()
    assert room["status"] == "critical"
